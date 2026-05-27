#!/usr/bin/env python3
# -*- coding: utf-8 -*-
# type: ignore[import]
from __future__ import annotations
"""
参数适用性矩阵生成程序

功能：
  双击运行，弹出文件夹选择对话框，选择利率装配源表文件夹后自动转换。
  输出到脚本同级 output/ 文件夹。
  转换成功弹窗提示，格式不一致弹窗报错。

源表结构约定（固定）：
  Row  2  : 产品名称（合并单元格，左填）
  Row  7  : 事件类型（正常 / 逾期 / 挪用），位于启用标识列上方
  Row 11  : 表头行，含“启用标识”列
  Row 12+ : 数据行
    Col 1 : 一级分类
    Col 2 : 二级分类
    Col 4 : 字段名称

转换前校验：
  - Row 11 Col 4 必须为“字段名称”
  - 每个 sheet 必须有 ≥1 个启用标识列
  - 启用标识列的 Row 7 必须能识别出事件类型（正常/逾期/挪用）

矩阵逻辑：
  - 参数唯一键 = 一级分类 + 二级分类 + 字段名称
  - 空分类按所有源表多数关系补齐
  - 按事件拆行：正常 / 逾期 / 挪用 / 未启用
  - 单元格：✅ 启用 / ❌ 未启用 / — 源产品未覆盖
  - 一级/二级分类输出时合并单元格

用法：
  python3 generate_param_applicability.py
  （或双击运行，选择文件夹即可）
"""

from __future__ import annotations

import re
import sys
from collections import Counter, OrderedDict, defaultdict
from pathlib import Path
from zipfile import BadZipFile

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

# ── 常量 ──────────────────────────────────────────────────────────────────────

EVENTS = ["正常", "逾期", "挪用"]
PRODUCT_ROW = 2
EVENT_ROW = 7
HEADER_ROW = 11
FIRST_PARAM_ROW = 12
L1_COL = 1
L2_COL = 2
FIELD_COL = 4

CHECK = "✅"
CROSS = "❌"
DASH = "—"

# ── 工具函数 ──────────────────────────────────────────────────────────────────

def norm(value) -> str:
    """去空白、换行,用于比较和 key。"""
    if value is None:
        return ""
    return re.sub(r"\s+", "", str(value).strip().replace("\n", " "))


def clean(value) -> str:
    """保留空格但合并连续空格,用于显示。"""
    if value is None:
        return ""
    return re.sub(r"\s+", " ", str(value).strip().replace("\n", " "))


def event_name(value) -> str:
    """从 Row 7 单元格值识别事件类型。"""
    text = norm(value)
    if "正常" in text:
        return "正常"
    if "逾期" in text or "本金违约" in text or text == "违约":
        return "逾期"
    if "挪用" in text:
        return "挪用"
    return ""


def is_enabled(value) -> bool:
    """判断启用标识值:1-启用 → True, 0-禁用 → False。"""
    if value is None:
        return False
    s = str(value).strip()
    return s.startswith("1") or "是" in s or "启用" in s or s.upper() == "TRUE"


def left_fill(ws, row: int, col: int) -> str:
    """返回当前单元格值,若为空则向左查找最近非空值(处理合并单元格)。"""
    for c in range(col, 0, -1):
        value = ws.cell(row, c).value
        if value not in (None, ""):
            return clean(value)
    return ""


def normalize_field_name(value) -> str:
    """标准化已知的字段名称变体,避免格式差异导致重复行。"""
    text = clean(value)
    aliases = {
        "当日还款计息标志（信贷使用）": "当日还款计息标志",
        "期限靠档规则（借据） 贷款": "期限靠档规则（借据）贷款",
    }
    return aliases.get(text, text)


def param_key(l1: str, l2: str, field_name: str) -> tuple[str, str, str]:
    return norm(l1), norm(l2), norm(normalize_field_name(field_name))


# ── 第一步:格式校验 ─────────────────────────────────────────────────────────

def validate_source_files(source_dir: Path, exclude_keyword: str) -> dict:
    """
    校验所有源表格式是否一致。
    返回: {
        "valid": bool,
        "files": [校验通过的文件路径],
        "errors": [(文件名, sheet名, 错误描述)],
        "warnings": [(文件名, sheet名, 警告描述)],
    }
    """
    errors = []
    warnings = []
    valid_files = []
    skipped = []

    xlsx_files = sorted(source_dir.glob("*.xlsx"))
    if not xlsx_files:
        return {"valid": False, "files": [], "errors": [("N/A", "N/A", "源文件夹中无 .xlsx 文件")], "warnings": []}

    for path in xlsx_files:
        fname = path.name

        # 跳过临时文件
        if fname.startswith("~$"):
            skipped.append((fname, "", "临时文件,已跳过"))
            continue

        # 排除关键字
        if exclude_keyword and exclude_keyword in fname:
            skipped.append((fname, "", f"文件名含排除关键字「{exclude_keyword}」,已跳过"))
            continue

        # 打开文件
        try:
            wb = load_workbook(path, data_only=True)
        except BadZipFile:
            errors.append((fname, "", "不是有效的 .xlsx 文件(损坏或非 xlsx 格式)"))
            continue
        except Exception as exc:
            errors.append((fname, "", f"打开失败: {type(exc).__name__}: {exc}"))
            continue

        file_ok = True
        for ws in wb.worksheets:
            # ── 基本结构检查 ──
            max_row = ws.max_row or 0
            max_col = ws.max_column or 0

            if max_row < FIRST_PARAM_ROW:
                warnings.append((fname, ws.title, f"行数不足({max_row} < {FIRST_PARAM_ROW}),跳过此 sheet"))
                continue

            if max_col < 10:
                warnings.append((fname, ws.title, f"列数不足({max_col} < 10),跳过此 sheet"))
                continue

            # ── Row 11 Col 4 = 字段名称 ──
            r11_d = ws.cell(HEADER_ROW, FIELD_COL).value
            if r11_d is None or "字段名称" not in str(r11_d):
                errors.append((fname, ws.title, f"Row {HEADER_ROW} Col {FIELD_COL} 应含「字段名称」,实际值: {repr(r11_d)}"))
                file_ok = False
                continue

            # ── Row 11 启用标识列检查 ──
            enabled_cols = []
            for col in range(1, max_col + 1):
                hdr = ws.cell(HEADER_ROW, col).value
                if hdr and "启用标识" in str(hdr):
                    enabled_cols.append(col)

            if not enabled_cols:
                errors.append((fname, ws.title, "Row 11 中未找到任何含「启用标识」的列"))
                file_ok = False
                continue

            # ── 启用标识列的 Row 7 必须能识别事件 ──
            unrecognized_events = []
            for col in enabled_cols:
                evt_raw = ws.cell(EVENT_ROW, col).value
                evt = event_name(evt_raw)
                if not evt:
                    unrecognized_events.append(f"Col {col}: Row 7 = {repr(evt_raw)}")

            if unrecognized_events:
                errors.append((
                    fname, ws.title,
                    f"以下启用标识列的 Row 7 无法识别事件类型(应含 正常/逾期/挪用):\n"
                    + "\n".join(f"  {e}" for e in unrecognized_events)
                ))
                file_ok = False
                continue

            # ── Row 2 产品名称检查 ──
            products_found = set()
            for col in enabled_cols:
                product = left_fill(ws, PRODUCT_ROW, col)
                if not product:
                    warnings.append((fname, ws.title, f"Col {col} 的 Row 2 产品名称为空(左填后仍为空)"))
                else:
                    products_found.add(product)

            # ── Row 12 首行数据检查 ──
            first_field = ws.cell(FIRST_PARAM_ROW, FIELD_COL).value
            if first_field is None or norm(first_field) == "":
                warnings.append((fname, ws.title, f"Row {FIRST_PARAM_ROW} Col {FIELD_COL} 字段名称为空"))

            if not file_ok:
                continue

        if file_ok:
            valid_files.append(path)

        wb.close()

    return {
        "valid": len(errors) == 0,
        "files": valid_files,
        "errors": errors,
        "warnings": warnings,
        "skipped": skipped,
    }


# ── 第二步:读取源表数据 ──────────────────────────────────────────────────────

def discover_source_sheets(valid_files: list[Path]):
    """读取所有校验通过的文件,提取参数数据。"""
    sheets = []
    included = []

    for path in valid_files:
        wb = load_workbook(path, data_only=True)
        included.append(path.name)

        for ws in wb.worksheets:
            max_row = ws.max_row or 0
            max_col = ws.max_column or 0
            if max_row < FIRST_PARAM_ROW or max_col < 10:
                continue

            # 找启用标识列
            enabled_cols: list[tuple[int, str, str]] = []
            for col in range(1, max_col + 1):
                hdr = ws.cell(HEADER_ROW, col).value
                evt = event_name(ws.cell(EVENT_ROW, col).value)
                if hdr and "启用标识" in str(hdr) and evt in EVENTS:
                    product = left_fill(ws, PRODUCT_ROW, col) or ws.title
                    enabled_cols.append((col, evt, product))

            if not enabled_cols:
                continue

            # 读取数据行
            rows = []
            last_l1 = ""
            last_l2 = ""
            for row in range(FIRST_PARAM_ROW, max_row + 1):
                raw_l1 = clean(ws.cell(row, L1_COL).value)
                raw_l2 = clean(ws.cell(row, L2_COL).value)
                if raw_l1:
                    last_l1 = raw_l1
                if raw_l2:
                    last_l2 = raw_l2

                raw_field = ws.cell(row, FIELD_COL).value
                field_name = normalize_field_name(raw_field)
                if not field_name or norm(field_name) == "字段名称":
                    continue

                enabled_values = {
                    (event, product, col): is_enabled(ws.cell(row, col).value)
                    for col, event, product in enabled_cols
                }
                rows.append((row, last_l1, last_l2, field_name, enabled_values))

            sheets.append((path.name, ws.title, rows))

        wb.close()

    return sheets, included


# ── 第三步:多数分类映射 ──────────────────────────────────────────────────────

def build_majority_maps(sheets):
    """从所有源表学习分类补齐规则。"""
    l2_to_l1 = defaultdict(Counter)
    field_l2_to_l1 = defaultdict(Counter)
    field_to_l1 = defaultdict(Counter)
    field_to_l2 = defaultdict(Counter)

    for _file, _sheet, rows in sheets:
        for _row, l1, l2, field_name, _values in rows:
            if l1 and l2:
                l2_to_l1[norm(l2)][l1] += 1
                field_l2_to_l1[(norm(field_name), norm(l2))][l1] += 1
            if l1:
                field_to_l1[norm(field_name)][l1] += 1
            if l2:
                field_to_l2[norm(field_name)][l2] += 1

    def majority(counter_map):
        return {key: counter.most_common(1)[0][0] for key, counter in counter_map.items() if counter}

    return {
        "l2_to_l1_counter": l2_to_l1,
        "l2_to_l1": majority(l2_to_l1),
        "field_l2_to_l1": majority(field_l2_to_l1),
        "field_to_l1": majority(field_to_l1),
        "field_to_l2": majority(field_to_l2),
    }


# ── 第四步:构建矩阵数据 ──────────────────────────────────────────────────────

def build_matrix_data(sheets, maps):
    products = OrderedDict()
    values = defaultdict(lambda: defaultdict(list))
    meta = {}
    source_detail = []
    fill_audit = []
    order = 0

    for file_name, sheet_name, rows in sheets:
        for row_no, raw_l1, raw_l2, field_name, enabled_values in rows:
            l1, l2 = raw_l1, raw_l2
            fill_reason = ""

            # 二级分类补齐
            if not l2:
                candidate = maps["field_to_l2"].get(norm(field_name), "")
                if candidate:
                    l2 = candidate
                    fill_reason += f"二级分类空,按字段名称多数补齐为「{l2}」;"

            # 一级分类补齐
            if not l1:
                candidate = (
                    maps["field_l2_to_l1"].get((norm(field_name), norm(l2)), "")
                    or maps["l2_to_l1"].get(norm(l2), "")
                    or maps["field_to_l1"].get(norm(field_name), "")
                )
                if candidate:
                    l1 = candidate
                    fill_reason += f"一级分类空,按多数关系补齐为「{l1}」;"

            if fill_reason:
                fill_audit.append([file_name, sheet_name, row_no, field_name, raw_l1, raw_l2, l1, l2, fill_reason])

            pk = param_key(l1, l2, field_name)
            if pk not in meta:
                meta[pk] = {"l1": l1, "l2": l2, "field_name": field_name, "first": order}
            order += 1

            for (event, product, col), enabled in enabled_values.items():
                products.setdefault(product, None)
                values[(pk, product)][event].append(enabled)
                source_detail.append([
                    field_name, l1, l2, product, event,
                    "启用" if enabled else "禁用",
                    file_name, sheet_name, col, row_no, fill_reason,
                ])

    product_list = list(products.keys())
    param_keys = sorted(meta.keys(), key=lambda key: meta[key]["first"])
    return products, product_list, values, meta, param_keys, source_detail, fill_audit


# ── 第五步:生成输出行 ────────────────────────────────────────────────────────

def event_cell(values, pk, product: str, event: str) -> str:
    seen = values.get((pk, product), {}).get(event, [])
    if not seen:
        return DASH
    return CHECK if any(seen) else CROSS


def any_seen(values, pk, product: str) -> bool:
    evmap = values.get((pk, product), {})
    return any(evmap.get(event) for event in EVENTS)


def any_enabled_for_event(values, product_list, pk, event: str) -> bool:
    return any(event_cell(values, pk, product, event) == CHECK for product in product_list)


def build_output_rows(values, meta, param_keys, product_list):
    rows = []
    for pk in param_keys:
        m = meta[pk]
        enabled_events = [event for event in EVENTS if any_enabled_for_event(values, product_list, pk, event)]
        if enabled_events:
            for event in enabled_events:
                rows.append([m["l1"], m["l2"], m["field_name"], event, pk])
        else:
            rows.append([m["l1"], m["l2"], m["field_name"], "未启用", pk])
    return rows


# ── 第六步:写入 Excel ────────────────────────────────────────────────────────

def style_header(row, fill):
    for cell in row:
        cell.font = Font(bold=True)
        cell.fill = fill
        cell.alignment = Alignment(wrap_text=True, horizontal="center", vertical="center")


def write_outputs(
    output_dir: Path,
    base_name: str,
    product_list,
    values,
    rows,
    source_detail,
    fill_audit,
    maps,
    included,
    validation_result,
):
    output_dir.mkdir(parents=True, exist_ok=True)
    matrix_path = output_dir / f"{base_name}.xlsx"
    log_path = output_dir / f"{base_name}_核对记录.xlsx"

    header_fill = PatternFill("solid", fgColor="D9EAF7")
    thin = Side(style="thin", color="D9D9D9")

    # ── 矩阵文件 ──
    wb = Workbook()
    ws = wb.active
    ws.title = "参数适用性矩阵"
    ws.append(["参数适用性矩阵 -- 产品×参数 适用性"])
    ws.append([
        "✅ = 启用 | ❌ = 不启用/不适用 | - = 源产品未覆盖该参数;"
        "空一级/二级分类按所有源表多数分类关系补齐;"
        "参数唯一键 = 一级分类 + 二级分类 + 字段名称。"
    ])
    ws.append([])
    ws.append(["一级分类", "二级分类", "字段名称", "适用方式"] + product_list)
    ws[1][0].font = Font(bold=True, size=14)
    style_header(ws[4], header_fill)

    for l1, l2, field_name, method, pk in rows:
        outrow = [l1, l2, field_name, method]
        for product in product_list:
            if method == "未启用":
                cell_value = (CROSS if any_seen(values, pk, product) else DASH)
            else:
                cell_value = event_cell(values, pk, product, method)
            outrow.append(cell_value)
        ws.append(outrow)

    last_data_row = 4 + len(rows)

    # 合并一级/二级分类
    def merge_same(col: int):
        row = 5
        while row <= last_data_row:
            value = ws.cell(row, col).value
            end = row
            while end + 1 <= last_data_row and ws.cell(end + 1, col).value == value:
                end += 1
            if value and end > row:
                ws.merge_cells(start_row=row, start_column=col, end_row=end, end_column=col)
                ws.cell(row, col).alignment = Alignment(vertical="center", horizontal="center", wrap_text=True)
            row = end + 1

    merge_same(1)
    merge_same(2)

    # 统计行
    ws.append([])
    stat_row = ws.max_row + 1
    for offset, (label, symbol) in enumerate([
        ("✅ 启用行数", CHECK),
        ("❌ 不启用/不适用行数", CROSS),
        ("- 未覆盖行数", DASH),
    ]):
        ws.cell(stat_row + offset, 1).value = "统计" if offset == 0 else ""
        ws.cell(stat_row + offset, 2).value = label
        for col in range(5, 5 + len(product_list)):
            vals = [ws.cell(r, col).value for r in range(5, last_data_row + 1)]
            ws.cell(stat_row + offset, col).value = sum(1 for v in vals if v == symbol)

    # 格式
    ws.freeze_panes = "E5"
    for col, width in {"A": 30, "B": 32, "C": 32, "D": 12}.items():
        ws.column_dimensions[col].width = width
    for col in range(5, 5 + len(product_list)):
        ws.column_dimensions[ws.cell(4, col).column_letter].width = 22
    for row in ws.iter_rows(min_row=1, max_row=ws.max_row, max_col=4 + len(product_list)):
        for cell in row:
            cell.alignment = Alignment(
                wrap_text=True, vertical="center",
                horizontal="center" if cell.column >= 4 else "left",
            )
            if cell.row >= 4:
                cell.border = Border(top=thin, bottom=thin, left=thin, right=thin)

    wb.save(matrix_path)

    # ── 核对记录文件 ──
    log_wb = Workbook()

    # 汇总
    summary = log_wb.active
    summary.title = "汇总"
    summary.append(["项目", "数量"])
    summary_rows = [
        ["参与源文件数", len(set(included))],
        ["源产品列数量", len(product_list)],
        ["矩阵行数(含适用方式拆行)", len(rows)],
        ["分类补齐行数", len(fill_audit)],
        ["来源明细行数", len(source_detail)],
    ]
    summary_rows.extend([
        [f"{method}行数", sum(1 for r in rows if r[3] == method)]
        for method in EVENTS + ["未启用"]
    ])
    for row in summary_rows:
        summary.append(row)
    style_header(summary[1], header_fill)
    summary.column_dimensions["A"].width = 34
    summary.column_dimensions["B"].width = 14

    # 分类补齐审计
    fill_ws = log_wb.create_sheet("分类补齐审计")
    fill_ws.append(["源文件", "工作表", "源行号", "字段名称", "原一级分类", "原二级分类", "补齐后一级分类", "补齐后二级分类", "补齐原因"])
    for row in fill_audit:
        fill_ws.append(row)
    style_header(fill_ws[1], header_fill)
    for col, width in zip("ABCDEFGHI", [70, 32, 10, 32, 28, 30, 28, 30, 80]):
        fill_ws.column_dimensions[col].width = width

    # 多数分类映射
    map_ws = log_wb.create_sheet("多数分类映射")
    map_ws.append(["映射类型", "键", "多数分类", "计数明细"])
    for key, counter in sorted(maps["l2_to_l1_counter"].items()):
        map_ws.append(["二级->一级", key, maps["l2_to_l1"].get(key, ""), "; ".join(f"{k}:{v}" for k, v in counter.most_common())])
    style_header(map_ws[1], header_fill)
    for col, width in {"A": 16, "B": 40, "C": 30, "D": 100}.items():
        map_ws.column_dimensions[col].width = width

    # 格式校验结果
    val_ws = log_wb.create_sheet("格式校验")
    val_ws.append(["类型", "文件", "工作表", "详情"])
    for fname, sheet, msg in validation_result.get("errors", []):
        val_ws.append(["❌ 错误", fname, sheet, msg])
    for fname, sheet, msg in validation_result.get("warnings", []):
        val_ws.append(["⚠️ 警告", fname, sheet, msg])
    for fname, sheet, msg in validation_result.get("skipped", []):
        val_ws.append(["⏭️ 跳过", fname, sheet, msg])
    style_header(val_ws[1], header_fill)
    for col, width in {"A": 12, "B": 70, "C": 32, "D": 100}.items():
        val_ws.column_dimensions[col].width = width

    # 全量来源明细
    source_ws = log_wb.create_sheet("全量来源明细")
    source_ws.append(["字段名称", "一级分类", "二级分类", "源产品", "适用方式", "启用标识", "源文件", "工作表", "启用标识列", "源行号", "分类补齐说明"])
    for row in source_detail:
        source_ws.append(row)
    style_header(source_ws[1], header_fill)

    log_wb.save(log_path)
    return matrix_path, log_path


# ── 主流程 ────────────────────────────────────────────────────────────────────

def run_conversion(source_dir: Path, output_dir: Path) -> tuple[bool, str]:
    """
    Execute the full conversion pipeline.
    Returns (success: bool, message: str).
    """
    BASE_NAME = "参数适用性_最终矩阵"
    EXCLUDE_KEYWORD = "行社"

    # Step 1: Validate format
    validation = validate_source_files(source_dir, EXCLUDE_KEYWORD)

    if not validation["valid"]:
        error_details = "\n".join(
            f"  [{fname}] {sheet}: {msg}"
            for fname, sheet, msg in validation["errors"]
        )
        return False, f"表文件格式不一致\n\n共 {len(validation['errors'])} 个错误:\n{error_details}"

    # Step 2: Read source sheets
    sheets, included = discover_source_sheets(validation["files"])

    # Step 3: Build majority fill maps
    maps = build_majority_maps(sheets)

    # Step 4: Build matrix data
    products, product_list, values, meta, param_keys, source_detail, fill_audit = build_matrix_data(sheets, maps)

    # Step 5: Generate output rows
    output_rows = build_output_rows(values, meta, param_keys, product_list)

    # Step 6: Write Excel
    matrix_path, log_path = write_outputs(
        output_dir=output_dir,
        base_name=BASE_NAME,
        product_list=product_list,
        values=values,
        rows=output_rows,
        source_detail=source_detail,
        fill_audit=fill_audit,
        maps=maps,
        included=included,
        validation_result=validation,
    )

    event_counts = Counter(r[3] for r in output_rows)
    detail_lines = [
        f"  参与源文件: {len(set(included))}",
        f"  产品列: {len(product_list)}",
        f"  参数唯一键: {len(param_keys)}",
        f"  矩阵行数: {len(output_rows)}",
        f"    正常: {event_counts.get('正常', 0)}  逾期: {event_counts.get('逾期', 0)}  挪用: {event_counts.get('挪用', 0)}  未启用: {event_counts.get('未启用', 0)}",
        f"  分类补齐行: {len(fill_audit)}",
        f"",
        f"  矩阵文件: {matrix_path}",
        f"  核对记录: {log_path}",
    ]
    return True, "转换成功\n\n" + "\n".join(detail_lines)


def main():
    import tkinter as tk
    from tkinter import filedialog, messagebox

    root = tk.Tk()
    root.withdraw()

    # Open folder picker dialog
    source_dir_str = filedialog.askdirectory(title="请选择利率装配源表文件夹")
    if not source_dir_str:
        sys.exit(0)

    source_dir = Path(source_dir_str).resolve()

    # Output dir: script_dir/output/
    script_dir = Path(__file__).resolve().parent
    output_dir = script_dir / "output"
    output_dir.mkdir(parents=True, exist_ok=True)

    # Run conversion
    try:
        success, message = run_conversion(source_dir, output_dir)
    except Exception as exc:
        messagebox.showerror("转换失败", f"发生异常:\n{type(exc).__name__}: {exc}")
        sys.exit(1)

    if success:
        messagebox.showinfo("完成", message)
    else:
        messagebox.showerror("转换失败", message)

    root.destroy()


if __name__ == "__main__":
    main()
